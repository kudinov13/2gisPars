import logging

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models import Company

logger = logging.getLogger(__name__)


def generate_excel_report(companies: list[Company], output_file: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Компании"
    headers = [
        "№", "Название", "Город / регион", "Адрес", "Телефон", "Сайт", "Есть сайт", "Есть мессенджер",
        "VK", "Telegram", "WhatsApp", "Рейтинг", "Количество отзывов", "Оценка лида", "Причины оценки",
        "Размер бизнеса", "Причины классификации", "Статус продаж", "Комментарий", "Возраст", "Молодая?",
        "Источники возраста", "Аудит сайта", "Достоверность аудита", "HTTP", "Время загрузки (с)",
        "Проблемы сайта (критичные)", "Проблемы сайта (минорные)", "Рубрики",
    ]
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    thin_border = Border(**{side: Side(style="thin") for side in ("left", "right", "top", "bottom")})
    for column, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=column, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    young_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    critical_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    lead_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
    for index, company in enumerate(companies, 1):
        values = [
            index, company.name, company.city, company.address, company.primary_phone,
            company.website if company.website and company.has_site else "нет сайта",
            "Да" if company.has_site else "Нет", "Да" if company.has_messenger else "Нет",
            company.vk_url or "-", company.telegram_url or "-", company.whatsapp_url or "-",
            company.rating, company.reviews_count, company.lead_score, "\n".join(company.lead_reasons),
            company.business_size, "\n".join(company.business_size_reasons), company.sales_status,
            company.sales_comment, company.age_display, "Да" if company.is_young else "-",
            _age_sources(company), company.site_audit_score, company.site_audit_confidence or "-",
            company.site_status_code, company.site_load_time,
            "\n".join(company.site_critical_issues) if company.site_critical_issues else "-",
            "\n".join(company.site_issues) if company.site_issues else "-", ", ".join(company.rubrics),
        ]
        row = index + 1
        for column, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if company.lead_score >= 70:
            for column in range(1, len(headers) + 1):
                ws.cell(row=row, column=column).fill = lead_fill
        elif company.is_young:
            for column in range(1, len(headers) + 1):
                ws.cell(row=row, column=column).fill = young_fill
        if company.site_critical_issues:
            ws.cell(row=row, column=27).fill = critical_fill

    widths = [5, 28, 18, 32, 18, 28, 11, 14, 24, 22, 22, 10, 14, 12, 35, 14, 32, 16, 30, 12, 10, 24, 12, 15, 9, 14, 38, 38, 28]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 40
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(companies) + 1}"
    ws.freeze_panes = "A2"
    wb.save(output_file)
    logger.info(f"Отчёт сохранён: {output_file}")


def _age_sources(company: Company) -> str:
    sources = []
    if company.domain_created_date:
        sources.append(f"домен: {company.domain_created_date.strftime('%Y-%m-%d')}")
    if company.vk_group_created_date:
        sources.append(f"VK: {company.vk_group_created_date.strftime('%Y-%m-%d')}")
    if company.first_review_date:
        sources.append(f"2GIS отзыв: {company.first_review_date.strftime('%Y-%m-%d')}")
    if company.age_unknown:
        sources.append("источник не найден")
    return "\n".join(sources) if sources else "-"
